# везде оставила комментарии, где что-то меняла/добавляла
# закоменченные строчки кода - это старый код
# оставила его, чтобы тебе было лучше видно разницу (до/после)
from tkinter import *
from tkinter import font
from tkinter import filedialog
from tkinter.messagebox import showerror, showinfo, askyesno

from CTkMessagebox import CTkMessagebox

import openpyxl
import docx

from src.PDF import PdfConvertScript as pdf


class Window:
    # начальные параметры окна
    # добавила строку
    def __init__(self, width=1100, height=550, title="Теория игр", resizable=(False, False)):
        self.root = Tk()
        self.root.title(title)
        self.root.geometry(f"{width}x{height}+70+30")
        self.root.resizable(resizable[0], resizable[1])
        self.root["background"] = "#1e1e1e"
        self.group_list = {1: [], 2: [], 3: [], 4: [], 5: [], 6: []}
        self.frame_parent = Frame
        self.frame_group = Frame
        self.count = 1

        # вот это добавила
        self.canvas = Canvas
        self.password = "123"  # пароль для входа
        self.entry = Entry
        self.entry_text = StringVar()
        self.frame_in_parent = Frame

    # прорисовка
    # тут тоже поменяла, там теперь другая функция
    def run(self):
        # self.draw_widgets()
        self.system_in()
        self.root.mainloop()

    # новая функция, в ней весь интерфейс для входа с паролем
    def system_in(self):
        _font_in = font.Font(family="Bookman Old Style", size=18, weight="normal", slant="roman", underline=True)
        _font_input = font.Font(family="Bookman Old Style", size=12, weight="normal", slant="roman")
        _font_button_in = font.Font(family="Bookman Old Style", size=14, weight="normal", slant="roman", underline=True)

        self.frame_in_parent = Frame(self.root, bg="#9400d3")
        self.frame_in_parent.pack(pady=(130, 0))

        (Label(self.frame_in_parent, text=' Программа для составления заданий \n по дисциплине: "ТЕОРИЯ ИГР" ',
               font=_font_in,
               fg="#bdbdbd", bg="#1e1e1e").pack(fill='x', pady=(20, 5), ipadx=5, ipady=3, padx=5))
        (Label(self.frame_in_parent, text="Введите пароль для доступа к генерации:", font=_font_input,
               bg="#1e1e1e", fg="#bdbdbd", width=45).pack(fill='x', pady=(5, 10), padx=5))

        frame_in_child = Frame(self.frame_in_parent, bg="#1e1e1e")
        frame_in_child.pack(fill='x', pady=(0, 10), padx=5, ipadx=7, ipady=3)

        frame_in_child.columnconfigure(index=0, weight=1)
        frame_in_child.columnconfigure(index=1, weight=1)
        frame_in_child.rowconfigure(index=0, weight=1)
        frame_in_child.rowconfigure(index=1, weight=1)

        def on_entry_click(event):
            if self.entry.get() == 'Введите пароль...':
                self.entry.delete(0, "end")
                self.entry.insert(0, '')
                self.entry.config(fg='#0a0a0a', show="☺")

        def on_focusout(event):
            if self.entry.get() == '':
                self.entry.insert(0, 'Введите пароль...')
                self.entry.config(fg='#6c6874', show="")

        # entry_text = StringVar()
        # show="☺"
        self.entry = Entry(frame_in_child, bg="#bdbdbd", width=28, font=_font_input, justify="center",
                           insertbackground="white", bd=0, textvariable=self.entry_text)

        self.entry.insert(0, 'Введите пароль...')
        self.entry.bind('<FocusIn>', on_entry_click)
        self.entry.bind('<FocusOut>', on_focusout)
        self.entry.config(fg='#6c6874')

        self.entry.pack(padx=10, pady=(15, 0), ipadx=7, ipady=3)

        enabled = IntVar()

        def pass_show():
            if enabled.get() == 1:
                text = self.entry_text.get()
                self.entry.config(show="")
                self.entry.delete(0, "end")
                self.entry.insert(0, text)
            else:
                self.entry.config(show="☺")

        (Checkbutton(frame_in_child, text="просмотр пароля", font=_font_input, bg="#1e1e1e", fg="#bdbdbd",
                     variable=enabled, command=pass_show).pack(padx=5, pady=3))
        (Button(frame_in_child, text="Войти", font=_font_button_in, bg="#9400d3", bd=0, fg="white", width=28, height=2,
                activebackground="#6c6874", activeforeground="#bdbdbd", command=self.sign_in).pack(pady=(5, 20)))

    # все виджеты групп
    # поменяла тут height на 310 у canvas
    def group_frame(self, parent, title_frame="Группа", count_group=None, count_student=None):
        # списки групп
        _font_group = font.Font(family="Bookman Old Style", size=14, weight="normal", slant="roman", underline=True)
        self.frame_group = LabelFrame(parent, text=f"Группа {self.count}", font=_font_group, fg="#b4b6c2",
                                      labelanchor="n", bg="#0a0a0a", bd=0, padx=5, pady=5)
        self.frame_group.grid(row=0, column=self.count - 1, padx=7, pady=4)

        # скролбарр
        canvas = Canvas(self.frame_group, bg="#ad90de", highlightthickness=0, height=310)
        scroll = Scrollbar(self.frame_group, orient="vertical", command=canvas.yview)
        canvas.config(width=300, yscrollcommand=scroll.set)
        scroll.grid(row=0, column=1, sticky='ns', padx=7, pady=4)
        canvas.grid(row=0, column=0, padx=7, pady=4)
        frame__ = Frame(canvas)
        canvas.create_window((0, 0), window=frame__)

        def conf(event):
            canvas.configure(scrollregion=canvas.bbox('all'))

        # строки внутри групп (ФИО)
        _font_label = font.Font(family="Bookman Old Style", size=11, weight="normal", slant="roman")
        for i in range(1, count_student + 1):
            label_group = Label(frame__, text=f"{i}. {self.group_list[self.count][i - 1]}", font=_font_label,
                                fg="#0a0a0a")
            label_group.pack(padx=4, fill="both")

        print(self.count)
        frame__.bind('<Configure>', conf)

    # чтение текста из документов
    # вот тут изменила чуток для txt файла
    def read_file(self):
        if (self.count <= len(self.group_list)):
            file_name = filedialog.askopenfilename()
            if file_name:
                if file_name.split('.')[1] == 'txt':
                    file = open(file_name, 'r', encoding='utf-8')
                    for line in file:
                        self.group_list[self.count].append(line.split('. ')[1].split('\n')[0])
                    # print(self.group_list)
                if file_name.split('.')[1] == 'xlsx' or file_name.split('.')[1] == 'xls':
                    book = openpyxl.open(file_name, read_only=True)
                    sheet = book.active
                    for row in range(1, sheet.max_row + 1):
                        self.group_list[self.count].append(f"{sheet[row][0].value}")
                    # print(self.group_list)
                if file_name.split('.')[1] == 'docx' or file_name.split('.')[1] == 'doc':
                    file = docx.Document(file_name)
                    for i in file.paragraphs:
                        line = i.text
                        self.group_list[self.count].append(line)
                    # print(self.group_list)

                self.group_frame(parent=self.frame_parent, count_group=self.count,
                                 count_student=len(self.group_list[self.count]))
                self.count += 1
        else:
            showerror(title="Превышено количество групп ☹", message="Очистите поле групп,\nчтобы добавить новые группы")

    # поменяла
    def draw_widgets(self):
        # вот тут перед canvas поставила везде self.
        self.canvas = Canvas(self.root, bg="#bdbdbd", highlightthickness=0)
        scroll = Scrollbar(self.canvas, orient="horizontal", command=self.canvas.xview)
        self.canvas.config(xscrollcommand=scroll.set)
        scroll.pack(side=BOTTOM, fill='x')
        self.canvas.pack(fill="both", expand=True, padx=10, pady=10)
        self.frame_parent = Frame(self.canvas, bg="#ad90de")
        self.canvas.create_window((0, 0), window=self.frame_parent, anchor='nw')

        def conf(event):
            self.canvas.configure(scrollregion=self.canvas.bbox('all'))

        self.frame_parent.bind('<Configure>', conf)

        # кнопки
        # чуток изменила в кнопках "Выбрать файл" и "Сгенерировать" параметры в .pack()
        # а еще изменила fg на 'white' в кнопках "Выбрать файл" и "Сгенерировать" (это сааамое важное!!)
        # ещё добавила новую кнопку "Очистить поле групп
        _font_button = font.Font(family="Bookman Old Style", size=12, weight="normal", slant="roman")

        (Button(self.root, text=f"Выбрать файл", font=_font_button, bg="#6c6874", bd=0, fg="white",
                activebackground="#9400d3", activeforeground="#ad90de", command=self.read_file).pack(pady=10))

        # не знаю как ты у себя оставил, но здесь я ее вынсла из функции в класс

        # def gener_tasks():
        #     print(f"Вот тут кнопочка генерации заданий работает")

        (Button(self.root, text="Сгенерировать", font=_font_button, bg="#9400d3", bd=0, fg="white", width=45, height=2,
                activebackground="#6c6874", activeforeground="#bdbdbd", command=self.gener_tasks).pack(anchor='s',
                                                                                                       padx=50))

        (Button(self.root, text="Очистить поле групп", font=_font_button, bg="#6c6874", bd=0, fg="#bdbdbd",
                activebackground="#9400d3", activeforeground="#ad90de", command=self.destroy_group).pack(pady=10))

    def destroy_group(self):
        self.group_list = {1: [], 2: [], 3: [], 4: [], 5: [], 6: []}
        self.count = 1
        self.frame_parent.destroy()

        self.frame_parent = Frame(self.canvas, bg="#ad90de")
        self.canvas.create_window((0, 0), window=self.frame_parent, anchor='nw')

        def conf(event):
            self.canvas.configure(scrollregion=self.canvas.bbox('all'))

        self.frame_parent.bind('<Configure>', conf)

    def showMessageInfo(self, message: str):
        CTkMessagebox(title="Info", message=message)

    def gener_tasks(self):
        groupCont = 0
        for key in self.group_list:
            if len(self.group_list[key]) != 0:
                message = pdf.CreatePDF(int(len(self.group_list[key])), list(self.group_list[key]), int(key)).title()
                groupCont += 1
                # self.showMessageInfo(message)
                # print(int(len(self.group_list[key])), list(self.group_list[key]))

        print(groupCont)
        if groupCont != 0:

            print(f"Вот тут кнопочка генерации заданий работает")
            result = askyesno(title="Предупреждение",
                              message="Возможна ошибка генерации pdf-файла"                                                         "\nПродолжить конвертацию?")
            if result:
                message = pdf.convertDocx2Pdf(groupCont)
                showinfo("Результат", message)
            else:
                showinfo("Результат", message=message)


    # тоже новенькая, проверяет верный ли пароль
    # если да, то появляется часть интерфейса со списками групп
    # если нет, то выводить сообщение об ошибке
    def sign_in(self):
        input_password = self.entry_text.get()
        if input_password == self.password:
            self.frame_in_parent.forget()
            self.draw_widgets()  # вот сюда засунула "старый" вызов функции из self.run()
        if input_password == "" or input_password == 'Введите пароль...':
            showinfo(title="Пустое поле ввода", message="Пароль не был введен, \nпожалуйста, введите пароль")
        elif input_password != self.password:
            showerror(title="Неверный пароль", message="Введен неверный пароль,\nВам отказано в доступе ☹")


if __name__ == "__main__":
    window = Window()
    window.run()
